# Micro-servico "orcamentos": preenche um PPTX com placeholders {{campo}}
# e converte em PDF via LibreOffice headless. Stateless: o PPTX chega em
# base64 no request (o n8n carrega o modelo do banco), o PDF volta em base64.
#
# Convencao de placeholders:
#   {{campo}}            -> substituido pelo valor; ausente vira "" (reportado)
#   {{opt:campo}}        -> igual, MAS se TODOS os opt: de uma forma (shape)
#                           ficarem vazios, a forma inteira e removida do
#                           slide (util para cards de opcao 2/3 nao usados)
#   {{opt-slide:campo}}  -> se o campo estiver vazio, o slide inteiro e
#                           removido (marcador pode ficar em qualquer texto
#                           do slide; o marcador em si sempre e apagado)

import base64
import os
import re
import shutil
import subprocess
import tempfile
import threading

from flask import Flask, jsonify, request
from pptx import Presentation

app = Flask(__name__)

TOKEN = os.environ.get("ORCAMENTOS_TOKEN", "")
PLACEHOLDER_RE = re.compile(r"\{\{\s*(opt:|opt-slide:)?\s*([A-Za-z0-9_.\-]+)\s*\}\}")

# LibreOffice nao lida bem com conversoes concorrentes no mesmo perfil;
# serializa e usa perfil isolado por chamada.
_lo_lock = threading.Lock()


def _iter_text_frames(shape):
    """Percorre recursivamente shapes (inclui grupos e tabelas) devolvendo
    todos os text_frames encontrados."""
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for sub in shape.shapes:
            yield from _iter_text_frames(sub)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def _shape_placeholders(shape):
    """Lista (kind, key) de todos os placeholders dentro de um shape."""
    found = []
    for tf in _iter_text_frames(shape):
        for para in tf.paragraphs:
            full = "".join(run.text for run in para.runs)
            for m in PLACEHOLDER_RE.finditer(full):
                found.append((m.group(1) or "", m.group(2)))
    return found


def _fill_paragraph(para, campos, relatorio):
    """Substitui placeholders no paragrafo inteiro (placeholders podem estar
    quebrados em varios runs). O texto final vai para o primeiro run — a
    formatacao do primeiro run vence dentro do paragrafo alterado."""
    full = "".join(run.text for run in para.runs)
    if "{{" not in full:
        return

    def sub(m):
        kind, key = (m.group(1) or ""), m.group(2)
        if kind == "opt-slide:":
            return ""  # marcador de slide: sempre some do texto
        val = campos.get(key)
        if val is None or str(val).strip() == "":
            if kind != "opt:":
                relatorio["vazios"].append(key)
            return ""
        relatorio["preenchidos"].append(key)
        return str(val)

    new = PLACEHOLDER_RE.sub(sub, full)
    if new != full and para.runs:
        para.runs[0].text = new
        for run in para.runs[1:]:
            run.text = ""


def _delete_slide(prs, slide):
    xml_slides = prs.slides._sldIdLst
    idx = prs.slides.index(slide)
    sld_id = list(xml_slides)[idx]
    prs.part.drop_rel(sld_id.rId)
    xml_slides.remove(sld_id)


def preencher(pptx_bytes, campos):
    prs = Presentation(_bytes_io(pptx_bytes))
    relatorio = {
        "preenchidos": [],
        "vazios": [],
        "slides_removidos": 0,
        "formas_removidas": 0,
    }

    def is_empty(key):
        val = campos.get(key)
        return val is None or str(val).strip() == ""

    # 1) remocao de slides marcados com {{opt-slide:campo}} vazio
    for slide in list(prs.slides):
        markers = [
            key
            for shape in slide.shapes
            for kind, key in _shape_placeholders(shape)
            if kind == "opt-slide:"
        ]
        if markers and all(is_empty(k) for k in markers):
            _delete_slide(prs, slide)
            relatorio["slides_removidos"] += 1

    # 2) remocao de formas cujos opt: ficaram todos vazios, depois preenchimento
    for slide in prs.slides:
        for shape in list(slide.shapes):
            opts = [key for kind, key in _shape_placeholders(shape) if kind == "opt:"]
            if opts and all(is_empty(k) for k in opts):
                shape._element.getparent().remove(shape._element)
                relatorio["formas_removidas"] += 1
                continue
            for tf in _iter_text_frames(shape):
                for para in tf.paragraphs:
                    _fill_paragraph(para, campos, relatorio)

    out = _bytes_io(None)
    prs.save(out)
    relatorio["vazios"] = sorted(set(relatorio["vazios"]))
    relatorio["preenchidos"] = sorted(set(relatorio["preenchidos"]))
    return out.getvalue(), relatorio


def _bytes_io(data):
    import io

    return io.BytesIO(data) if data is not None else io.BytesIO()


def pptx_para_pdf(pptx_bytes):
    workdir = tempfile.mkdtemp(prefix="orc-")
    try:
        src = os.path.join(workdir, "orcamento.pptx")
        with open(src, "wb") as f:
            f.write(pptx_bytes)
        profile = os.path.join(workdir, "lo-profile")
        cmd = [
            "soffice",
            "--headless",
            "--norestore",
            "-env:UserInstallation=file://" + profile,
            "--convert-to",
            "pdf:impress_pdf_Export",
            "--outdir",
            workdir,
            src,
        ]
        with _lo_lock:
            proc = subprocess.run(cmd, capture_output=True, timeout=180)
        pdf_path = os.path.join(workdir, "orcamento.pdf")
        if proc.returncode != 0 or not os.path.exists(pdf_path):
            raise RuntimeError(
                "LibreOffice falhou: " + proc.stderr.decode("utf-8", "ignore")[-500:]
            )
        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def _auth_ok(req):
    return TOKEN and req.headers.get("X-Orcamentos-Token") == TOKEN


@app.get("/health")
def health():
    return jsonify({"ok": True})


@app.post("/inspecionar")
def inspecionar():
    """Lista os placeholders de um PPTX — usado ao cadastrar um modelo novo."""
    if not _auth_ok(request):
        return jsonify({"ok": False, "erro": "nao autorizado"}), 401
    body = request.get_json(force=True)
    pptx_bytes = base64.b64decode(body["pptx_b64"])
    prs = Presentation(_bytes_io(pptx_bytes))
    campos = []
    for num, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            for kind, key in _shape_placeholders(shape):
                campos.append({"slide": num, "tipo": kind or "normal", "campo": key})
    return jsonify({"ok": True, "total_slides": len(prs.slides), "campos": campos})


@app.post("/render")
def render():
    """Preenche o modelo e devolve o PDF (ou PPTX) em base64.

    Body: {"pptx_b64": "...", "campos": {"cliente": "Dona Ana", ...},
           "formato": "pdf" | "pptx"}
    """
    if not _auth_ok(request):
        return jsonify({"ok": False, "erro": "nao autorizado"}), 401
    body = request.get_json(force=True)
    pptx_bytes = base64.b64decode(body["pptx_b64"])
    campos = body.get("campos") or {}
    formato = body.get("formato", "pdf")

    preenchido, relatorio = preencher(pptx_bytes, campos)
    if formato == "pptx":
        return jsonify(
            {
                "ok": True,
                "pptx_b64": base64.b64encode(preenchido).decode(),
                "relatorio": relatorio,
            }
        )
    pdf = pptx_para_pdf(preenchido)
    return jsonify(
        {
            "ok": True,
            "pdf_b64": base64.b64encode(pdf).decode(),
            "relatorio": relatorio,
        }
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8080")))
