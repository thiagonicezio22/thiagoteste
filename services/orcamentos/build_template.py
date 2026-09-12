#!/usr/bin/env python3
# Gera o modelo PPTX oficial de orcamento a partir de um layout spec JSON
# + imagens JPEG de pagina (fundo full-bleed). O resultado e o arquivo que:
#   1) o micro-servico (app.py) preenche via placeholders {{campo}} e
#      {{opt:campo}} e converte em PDF (LibreOffice headless);
#   2) o Thiago pode abrir e editar direto no PowerPoint.
#
# Convencoes (ver app.py / README.md):
#   - cada elemento de texto do layout vira UMA caixa de texto propria, sem
#     preenchimento nem borda, contendo o placeholder do campo;
#   - elementos com "if" no layout usam {{opt:campo}} — quando todos os opt:
#     de uma forma ficarem vazios, o app.py remove a forma inteira (cards de
#     opcao 2/3 nao usados somem do slide);
#   - os rects "ifEmpty" do layout NAO viram shapes: no PPTX quem some e o
#     proprio texto opt:, o fundo (imagem) permanece.
#
# Uso:
#   python3 build_template.py [layout.json] [saida.pptx]
#   Sem argumentos usa o modelo piloto ambiente-aquatico.

import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

# Slide 16:9 de 10 x 5.625 pol = 720 x 405 pt
SLIDE_W_PT = 720.0
SLIDE_H_PT = 405.0
EMU_PER_PT = 12700

# Fontes logicas do layout -> fontes reais no PPTX. No container do
# micro-servico o LibreOffice usa as clones metricas Carlito (=Calibri) e
# Caladea (=Cambria), entao o PDF mantem as mesmas medidas do PowerPoint.
FONT_MAP = {
    "serif-bold": ("Cambria", True),
    "sans": ("Calibri", False),
    "sans-bold": ("Calibri", True),
}

# Distancia aproximada do topo da caixa (margens zeradas) ate a baseline da
# primeira linha, em fracao do tamanho da fonte (metrica winAscent de
# Cambria/Caladea e Calibri/Carlito). Usada para ancorar o texto de forma
# que a baseline caia proxima da especificada no layout.
ASCENT = {
    "serif-bold": 0.950,
    "sans": 0.952,
    "sans-bold": 0.952,
}

# Altura da caixa de uma linha, em fracao do tamanho da fonte. Linha simples
# ocupa ~1.22em; a folga minima faz o autofit ("shrink") so agir quando o
# texto realmente nao cabe na largura e quebra para uma segunda linha.
LINE_BOX_FACTOR = 1.25


def _pt(v):
    """Converte pontos em EMU (unidade interna do PPTX)."""
    return Emu(int(round(v * EMU_PER_PT)))


def _rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#"))


def _placeholder(element, key):
    """Placeholder do campo; elementos condicionais ("if") usam opt: para a
    forma inteira sumir quando o campo vier vazio (convencao do app.py)."""
    prefix = "opt:" if element.get("if") else ""
    return "{{%s%s}}" % (prefix, key)


def _style_run(run, font_key, size_pt, hex_color):
    name, bold = FONT_MAP[font_key]
    font = run.font
    font.name = name
    font.bold = bold
    font.size = Pt(size_pt)
    font.color.rgb = _rgb(hex_color)


def _add_text_element(slide, element, sx, sy):
    """Cria a caixa de texto de um elemento text/wrap/line do layout."""
    etype = element["type"]
    align_right = element.get("align") == "right"
    max_width = element["maxWidth"]

    # tamanho/fonte de referencia para ancorar a baseline (em "line" a
    # baseline e compartilhada: vale o maior ascent entre os segments)
    if etype == "line":
        ref_size, ref_font = max(
            (seg["size"], seg["font"]) for seg in element["segments"]
        )
    else:
        ref_size, ref_font = element["size"], element["font"]

    # geometria: pt do layout -> pt do slide -> EMU. Em align right o "x" do
    # layout e a borda DIREITA do texto; a caixa se estende para a esquerda.
    left_pt = (element["x"] - max_width) if align_right else element["x"]
    top_pt = element["baseline"] - ASCENT[ref_font] * ref_size
    if etype == "wrap":
        height_pt = element["lineHeight"] * element["maxLines"] + 2
    else:
        height_pt = LINE_BOX_FACTOR * ref_size

    box = slide.shapes.add_textbox(
        _pt(left_pt * sx), _pt(top_pt * sy), _pt(max_width * sx), _pt(height_pt * sy)
    )
    box.fill.background()
    box.line.fill.background()

    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    # shrink do layout = "encolher fonte ate caber": autofit do PowerPoint.
    # A caixa tem altura de 1 linha; texto largo demais quebraria linha, e o
    # autofit reduz a fonte ate caber de novo. Em "wrap" o autofit protege
    # contra descricoes maiores que maxLines.
    if element.get("shrink") or etype == "wrap":
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    else:
        tf.auto_size = MSO_AUTO_SIZE.NONE

    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT if align_right else PP_ALIGN.LEFT
    if etype == "wrap":
        para.line_spacing = Pt(element["lineHeight"] * sy)

    if etype == "line":
        # um paragrafo so, um run por segment para preservar estilos mistos
        # (obs.: apos o preenchimento o app.py consolida o texto no primeiro
        # run — a formatacao do primeiro trecho vence, como documentado)
        for seg in element["segments"]:
            run = para.add_run()
            run.text = seg.get("text") or _placeholder(element, seg["key"])
            _style_run(run, seg["font"], seg["size"], seg["color"])
        first_key = next(s["key"] for s in element["segments"] if "key" in s)
        box.name = "linha_" + first_key
    else:
        run = para.add_run()
        run.text = _placeholder(element, element["key"])
        _style_run(run, element["font"], element["size"], element["color"])
        box.name = "txt_" + element["key"]

    return box


def build(layout_path, out_path):
    with open(layout_path, encoding="utf-8") as f:
        layout = json.load(f)

    # imagens de pagina ficam na pasta com o mesmo nome base do layout
    img_dir = layout_path
    for suffix in (".layout.json", ".json"):
        if img_dir.endswith(suffix):
            img_dir = img_dir[: -len(suffix)]
            break

    # fator de escala layout -> slide (pageH 405.014pt vs slide 405pt)
    sx = SLIDE_W_PT / layout["pageW"]
    sy = SLIDE_H_PT / layout["pageH"]

    prs = Presentation()
    prs.slide_width = _pt(SLIDE_W_PT)  # 10 pol
    prs.slide_height = _pt(SLIDE_H_PT)  # 5.625 pol
    blank = prs.slide_layouts[6]  # layout "Blank"

    for num, page in enumerate(layout["pages"], start=1):
        slide = prs.slides.add_slide(blank)

        # fundo full-bleed: a arte fixa da pagina ja vem rasterizada no JPEG
        img_path = os.path.join(img_dir, "page%02d.jpg" % (page["img"] + 1))
        picture = slide.shapes.add_picture(
            img_path, 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        picture.name = "fundo_pag%02d" % num

        for element in page.get("elements", []):
            if element["type"] == "rect":
                continue  # rects ifEmpty nao viram shapes (ver cabecalho)
            _add_text_element(slide, element, sx, sy)

    prs.save(out_path)
    total_el = sum(
        1
        for page in layout["pages"]
        for el in page.get("elements", [])
        if el["type"] != "rect"
    )
    print(
        "gerado %s: %d slides, %d caixas de placeholder"
        % (out_path, len(layout["pages"]), total_el)
    )


if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))
    layout_path = (
        sys.argv[1]
        if len(sys.argv) > 1
        else os.path.join(base, "templates", "ambiente-aquatico.layout.json")
    )
    out_path = (
        sys.argv[2]
        if len(sys.argv) > 2
        else os.path.join(base, "templates", "ambiente-aquatico.pptx")
    )
    build(layout_path, out_path)
